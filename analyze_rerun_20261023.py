#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""分析 2026-10-23 补跑结果（应审稿人 P4）：
1) 同日 runA vs runB 一致率（当前端点版本的运行间稳定性）
2) 10月 vs 7月9日 原始分（端点漂移 + 噪声）
3) 考官-LLM 分歧集进出条数（工作流相关性）
4) 10月分数对 818 条专家多数共识参考标准的准确率（指纹连接 data_release CSV）
"""
import json, glob, os, itertools, math
from pathlib import Path
import pandas as pd
import openpyxl

OUT = Path(os.environ.get("STT_OUTPUT_DIR", "../output"))
CSV = Path(os.environ.get("ITEM_CSV", "item_level_scores_anonymized.csv"))

def load_xlsx_items(path):
    ws = openpyxl.load_workbook(path).active
    hdr = [c.value for c in ws[1]]
    i_seq, i_max, i_h, i_ai = 0, hdr.index("分值"), hdr.index("人工得分"), hdr.index("AI得分")
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not isinstance(r[0], (int, float)):
            continue
        rows.append((int(r[0]), float(r[i_max]), r[i_h], r[i_ai]))
    return rows

def load_dir(d):
    out = {}
    for f in glob.glob(str(d / "*.xlsx")):
        name = os.path.basename(f)[:-5]
        out[name] = load_xlsx_items(f)
    return out

def wilson(k, n, z=1.96):
    if n == 0: return (0, 0)
    p = k / n
    d = 1 + z*z/n
    c = (p + z*z/(2*n)) / d
    h = z * math.sqrt(p*(1-p)/n + z*z/(4*n*n)) / d
    return (c - h, c + h)

july_main = load_dir(OUT / "AI_scores")              # 240 analysis, 2026-07-09
july_dev  = load_dir(OUT / "AI_scores" / "预实验")   # 36 dev, 2026-07-09 snapshot
runA      = load_dir(OUT / "rerun_20261023" / "A")
runB      = load_dir(OUT / "rerun_20261023" / "B")

analysis_names = sorted(set(july_main) & set(runA) & set(runB))
dev_names      = sorted(set(july_dev)  & set(runA) & set(runB))
print(f"encounters: analysis {len(analysis_names)}, dev {len(dev_names)}")

def paired(setA, setB, names, col_a=3, col_b=3):
    """item-level agreement between two score dicts (AI col idx 3)"""
    tot = agree = 0
    for n in names:
        a = {r[0]: r for r in setA[n]}
        b = {r[0]: r for r in setB[n]}
        for seq in a:
            if seq in b:
                va, vb = a[seq][col_a], b[seq][col_b]
                if va is None or vb is None: continue
                tot += 1
                if va == vb: agree += 1
    return agree, tot

def examiner_col(d, names):
    """{name: {seq: examiner}}"""
    return {n: {r[0]: r[2] for r in d[n]} for n in names}

report = {}

# ---------- 1) same-day A vs B ----------
for label, names in [("all 276", analysis_names + dev_names),
                     ("analysis 240", analysis_names), ("dev 36", dev_names)]:
    k, n = paired(runA, runB, names)
    lo, hi = wilson(k, n)
    report[f"sameday_AB_{label}"] = {"agree": k, "n": n, "pct": round(k/n*100, 2), "ci": [round(lo*100,2), round(hi*100,2)]}
    print(f"[same-day A vs B | {label}] {k}/{n} = {k/n*100:.2f}% (95% CI {lo*100:.2f}-{hi*100:.2f})")

# ---------- 2) October vs July ----------
for label, july, names in [("analysis 240", july_main, analysis_names), ("dev 36", july_dev, dev_names)]:
    for run_name, run in [("A", runA), ("B", runB)]:
        k, n = paired(july, run, names)
        report[f"drift_{run_name}_{label}"] = {"agree": k, "n": n, "pct": round(k/n*100, 2)}
        print(f"[Oct {run_name} vs Jul-09 | {label}] {k}/{n} = {k/n*100:.2f}%")

# ---------- 3) discordance-set flips (analysis set, examiner vs LLM) ----------
def flip_stats(run, names):
    july_disc = oct_disc = both = flip_in = flip_out = tot = 0
    for n in names:
        jmap = {r[0]: r for r in july_main[n]}
        omap = {r[0]: r for r in run[n]}
        for seq, jr in jmap.items():
            if seq not in omap: continue
            e, jai, oai = jr[2], jr[3], omap[seq][3]
            if e is None or jai is None or oai is None: continue
            tot += 1
            jd = (e != jai); od = (e != oai)
            july_disc += jd; oct_disc += od; both += (jd and od)
            flip_in  += (not jd) and od
            flip_out += jd and (not od)
    return {"tot": tot, "july_disc": july_disc, "oct_disc": oct_disc,
            "july_pct": round(july_disc/tot*100, 2), "oct_pct": round(oct_disc/tot*100, 2),
            "flip_in": flip_in, "flip_out": flip_out,
            "stable_discordant": both, "stable_concordant": tot - july_disc - oct_disc + both}
for run_name, run in [("A", runA), ("B", runB)]:
    s = flip_stats(run, analysis_names)
    report[f"flips_{run_name}"] = s
    print(f"[flips {run_name}] items={s['tot']}; discordant Jul={s['july_disc']} ({s['july_pct']}%) -> Oct={s['oct_disc']} ({s['oct_pct']}%); in={s['flip_in']}, out={s['flip_out']}, stable-discordant={s['stable_discordant']}")

# ---------- 4) accuracy vs 818 reference via fingerprint join ----------
df = pd.read_csv(CSV)
enc_groups = {e: g.sort_values("item_seq") for e, g in df.groupby("encounter_id")}
# fingerprint: (sp_case, tuple((examiner, v4pro) per item))
fp_index = {}
for e, g in enc_groups.items():
    fp = (g.sp_case.iloc[0], tuple(zip(g.examiner_score, g.deepseek_v4_pro)))
    fp_index.setdefault(fp, []).append(e)

name2enc = {}
ambiguous = []
for n in analysis_names:
    jrows = sorted(july_main[n])
    sp = n.split("SP")[1] if "SP" in n else None
    sp = "SP" + n[n.index("SP")+2:n.index("SP")+4] if "SP" in n else None
    sp = "SP0" + n[n.index("SP")+2] if (n[n.index("SP")+2].isdigit()) else sp
    m = n[n.index("SP"):n.index("SP")+4]  # e.g. SP04
    fp = (m, tuple((r[2], r[3]) for r in jrows))
    hits = fp_index.get(fp, [])
    if len(hits) == 1:
        name2enc[n] = hits[0]
    else:
        ambiguous.append((n, len(hits)))
print(f"fingerprint join: matched {len(name2enc)}/{len(analysis_names)}, ambiguous {len(ambiguous)}")
if ambiguous:
    print("  ambiguous:", ambiguous[:5])

# expert majority on 818
import collections
def majority(r):
    vs = [r.expert_wst, r.expert_fj, r.expert_gxx]
    vs = [v for v in vs if pd.notna(v)]
    c = collections.Counter(vs)
    if not c: return None
    v, k = c.most_common(1)[0]
    return v if k >= 2 else None

df["consensus"] = df.apply(majority, axis=1)
ref = df[df.consensus.notna()][["encounter_id", "item_seq", "examiner_score", "deepseek_v4_pro", "consensus"]]

acc = {}
for run_name, run in [("July", july_main), ("OctA", runA), ("OctB", runB)]:
    k = n = 0
    for name, e in name2enc.items():
        sub = ref[ref.encounter_id == e]
        if sub.empty: continue
        omap = {r[0]: r[3] for r in run[name]}
        for _, r in sub.iterrows():
            v = omap.get(int(r.item_seq))
            if v is None: continue
            n += 1
            if v == r.consensus: k += 1
    lo, hi = wilson(k, n)
    acc[run_name] = {"match": k, "n": n, "pct": round(k/n*100, 1), "ci": [round(lo*100,1), round(hi*100,1)]}
    print(f"[accuracy vs 818 reference | {run_name}] {k}/{n} = {k/n*100:.1f}% (95% CI {lo*100:.1f}-{hi*100:.1f})")
report["accuracy_vs_818"] = acc
report["fingerprint_matched"] = len(name2enc)

# meta
meta_p = OUT / "rerun_20261023" / "meta.json"
if meta_p.exists():
    report["meta"] = json.loads(meta_p.read_text())

out_path = OUT / "rerun_20261023" / "analysis_results.json"
out_path.write_text(json.dumps(report, ensure_ascii=False, indent=1))
print("saved:", out_path)
